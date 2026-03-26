"""
Data export utilities — Excel (.xlsx) and PDF.
"""

import io
from typing import List, Dict, Any, Optional
from openpyxl import Workbook


def export_excel(
    rows: List[Dict[str, Any]],
    columns: Optional[List[str]] = None,
    sheet_name: str = "Sheet1",
) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    if not rows:
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    cols = columns or list(rows[0].keys())

    for col_idx, col_name in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name.replace("_", " ").title())
        cell.font = cell.font.copy(bold=True)

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, col_name in enumerate(cols, start=1):
            ws.cell(row=row_idx, column=col_idx, value=row.get(col_name, ""))

    for col_cells in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col_cells)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max_len + 4, 50)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_pdf_simple(
    rows: List[Dict[str, Any]],
    columns: Optional[List[str]] = None,
    title: str = "Report",
) -> io.BytesIO:
    try:
        from weasyprint import HTML
    except ImportError:
        raise RuntimeError("weasyprint is not installed.")

    cols = columns or (list(rows[0].keys()) if rows else [])
    headers = "".join(f"<th>{c.replace('_', ' ').title()}</th>" for c in cols)
    body_rows = ""
    for row in rows:
        cells = "".join(f"<td>{row.get(c, '')}</td>" for c in cols)
        body_rows += f"<tr>{cells}</tr>"

    html_str = f"""<!DOCTYPE html><html><head><meta charset="utf-8">
    <style>body{{font-family:sans-serif;margin:20px}}h1{{font-size:18px}}
    table{{width:100%;border-collapse:collapse;font-size:12px}}
    th,td{{border:1px solid #ccc;padding:6px 8px;text-align:left}}
    th{{background:#f5f5f5}}</style></head><body>
    <h1>{title}</h1><table><thead><tr>{headers}</tr></thead>
    <tbody>{body_rows}</tbody></table></body></html>"""

    buf = io.BytesIO()
    HTML(string=html_str).write_pdf(buf)
    buf.seek(0)
    return buf
